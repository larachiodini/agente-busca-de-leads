"""
Ferramentas do Agente de Leads.
Cada função faz UMA coisa e pode ser testada sozinha:
    python ferramentas.py
"""

import csv
import os
import re

import httpx
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

load_dotenv()

NAVEGADOR = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/124.0 Safari/537.36"}

# DDDs válidos no Brasil (evita "telefones" falsos tirados de códigos de página)
DDDS_VALIDOS = {
    11, 12, 13, 14, 15, 16, 17, 18, 19, 21, 22, 24, 27, 28,
    31, 32, 33, 34, 35, 37, 38, 41, 42, 43, 44, 45, 46, 47, 48, 49,
    51, 53, 54, 55, 61, 62, 63, 64, 65, 66, 67, 68, 69,
    71, 73, 74, 75, 77, 79, 81, 82, 83, 84, 85, 86, 87, 88, 89,
    91, 92, 93, 94, 95, 96, 97, 98, 99,
}

PADRAO_TELEFONE = re.compile(
    r"0800[\s.-]?\d{3}[\s.-]?\d{4}"                      # 0800 xxx xxxx
    r"|\b[34]00\d[\s.-]?\d{4}\b"                         # 4004-xxxx / 3003-xxxx
    r"|(?:\+?55[\s.-]*)?\(?0?\d{2}\)?[\s.-]*\d{4,5}[\s.-]?\d{4}"  # (19) 3577-9705
)


# ----------------------------------------------------------------------
# 1. Telefones: encontrar, validar e padronizar
# ----------------------------------------------------------------------
def padronizar_telefone_fixo(texto):
    """Recebe um telefone em qualquer formato. Devolve '(19) 3577-9705'
    se for FIXO válido, ou None se for celular/inválido."""
    digitos = re.sub(r"\D", "", texto)

    if digitos.startswith("0800") and len(digitos) == 11:
        return f"0800 {digitos[4:7]} {digitos[7:]}"
    if len(digitos) == 8 and digitos[:3] in ("400", "300"):
        return f"{digitos[:4]}-{digitos[4:]}"

    if digitos.startswith("55") and len(digitos) in (12, 13):
        digitos = digitos[2:]
    if digitos.startswith("0") and len(digitos) == 11:
        digitos = digitos[1:]

    # Fixo = DDD (2) + 8 dígitos começando com 2, 3, 4 ou 5
    if len(digitos) != 10:
        return None
    ddd, numero = int(digitos[:2]), digitos[2:]
    if ddd not in DDDS_VALIDOS or numero[0] not in "2345":
        return None
    return f"({ddd}) {numero[:4]}-{numero[4:]}"


def extrair_telefones_fixos(texto):
    """Acha todos os telefones fixos válidos dentro de um texto, sem repetir."""
    encontrados = []
    for trecho in PADRAO_TELEFONE.findall(texto or ""):
        telefone = padronizar_telefone_fixo(trecho)
        if telefone and telefone not in encontrados:
            encontrados.append(telefone)
    return encontrados


def limpar_cnpj(cnpj):
    return re.sub(r"\D", "", str(cnpj)).zfill(14)


# ----------------------------------------------------------------------
# 2. Receita Federal (via BrasilAPI, gratuita)
# ----------------------------------------------------------------------
def consultar_cnpj(cnpj):
    """Dados oficiais da Receita: nome, situação, atividade e telefones."""
    cnpj = limpar_cnpj(cnpj)
    try:
        resposta = httpx.get(f"https://brasilapi.com.br/api/cnpj/v1/{cnpj}", timeout=20)
        resposta.raise_for_status()
        dados = resposta.json()
    except Exception as erro:
        return {"cnpj": cnpj, "erro": f"Não foi possível consultar: {erro}"}

    telefones_brutos = " ".join(
        str(dados.get(campo) or "") for campo in ("ddd_telefone_1", "ddd_telefone_2", "ddd_fax")
    )
    return {
        "cnpj": cnpj,
        "razao_social": dados.get("razao_social"),
        "nome_fantasia": dados.get("nome_fantasia"),
        "situacao": dados.get("descricao_situacao_cadastral"),
        "atividade_principal": dados.get("cnae_fiscal_descricao"),
        "cidade": f"{dados.get('municipio')}/{dados.get('uf')}",
        "telefones_fixos": extrair_telefones_fixos(telefones_brutos),
    }


# ----------------------------------------------------------------------
# 3. Busca no Google (via Serper)
# ----------------------------------------------------------------------
def buscar_google(consulta, quantidade=10):
    """Pesquisa no Google e devolve título, link, resumo e telefones vistos no resumo."""
    chave = os.getenv("SERPER_API_KEY")
    if not chave:
        return {"erro": "SERPER_API_KEY não encontrada no arquivo .env"}
    try:
        resposta = httpx.post(
            "https://google.serper.dev/search",
            headers={"X-API-KEY": chave, "Content-Type": "application/json"},
            json={"q": consulta, "gl": "br", "hl": "pt-br", "num": quantidade},
            timeout=20,
        )
        resposta.raise_for_status()
        dados = resposta.json()
    except Exception as erro:
        return {"erro": f"Falha na busca: {erro}"}

    resultados = []
    for item in dados.get("organic", []):
        texto = f"{item.get('title', '')} {item.get('snippet', '')}"
        resultados.append({
            "titulo": item.get("title"),
            "link": item.get("link"),
            "resumo": item.get("snippet"),
            "telefones_no_resumo": extrair_telefones_fixos(texto),
        })

    painel = dados.get("knowledgeGraph", {})
    return {
        "consulta": consulta,
        "painel_google": {
            "nome": painel.get("title"),
            "site": painel.get("website"),
            "telefones": extrair_telefones_fixos(str(painel)),
        } if painel else None,
        "resultados": resultados,
    }


# ----------------------------------------------------------------------
# 4. Ler uma página (site da empresa, CNPJ Biz etc.)
# ----------------------------------------------------------------------
PALAVRAS_CONTATO = ("contato", "contact", "fale", "atendimento", "sac", "onde-estamos", "unidades")


def ler_site(url):
    """Abre uma página e devolve: telefones fixos, links de contato e um trecho do texto."""
    if not url.startswith("http"):
        url = "https://" + url
    try:
        resposta = httpx.get(url, headers=NAVEGADOR, timeout=20, follow_redirects=True)
        resposta.raise_for_status()
        html = resposta.text
    except Exception as erro:
        return {"url": url, "erro": f"Não foi possível abrir: {erro}"}

    html_sem_codigo = re.sub(r"(?is)<(script|style|noscript)[^>]*>.*?</\1>", " ", html)
    texto = re.sub(r"(?s)<[^>]+>", " ", html_sem_codigo)
    texto = re.sub(r"\s+", " ", texto).strip()

    links = re.findall(r'href=["\']([^"\']+)["\']', html_sem_codigo)
    links_contato = []
    for link in links:
        if any(palavra in link.lower() for palavra in PALAVRAS_CONTATO):
            completo = str(httpx.URL(str(resposta.url)).join(link))
            if completo not in links_contato:
                links_contato.append(completo)

    return {
        "url": str(resposta.url),
        "telefones_fixos": extrair_telefones_fixos(html_sem_codigo),  # inclui links "tel:"
        "links_de_contato": links_contato[:10],
        "trecho_do_texto": texto[:3000],
    }


# ----------------------------------------------------------------------
# 5. Ler a lista exportada (Speedio, Driva ou outro CSV/XLSX)
# ----------------------------------------------------------------------
MAPA_COLUNAS = {
    "cnpj": ("cnpj",),
    "nome": ("razão social", "razao social", "razao_social", "nome", "empresa"),
    "telefone": ("telefone", "fone"),
    "site": ("site", "website", "url", "domínio", "dominio"),
    "setor": ("setor", "segmento"),
    "subsetor": ("subsetor", "sub-setor", "subsegmento", "cnae"),
    "faturamento": ("faturamento", "receita"),
}


def _achar_coluna(cabecalhos, opcoes, ja_usadas):
    for opcao in opcoes:
        for indice, nome in enumerate(cabecalhos):
            if indice not in ja_usadas and opcao in str(nome or "").strip().lower():
                return indice
    return None


def carregar_lista_speedio(caminho):
    """Lê o arquivo exportado do Speedio e devolve uma lista de empresas padronizada."""
    if caminho.lower().endswith(".xlsx"):
        planilha = load_workbook(caminho, read_only=True).active
        linhas = [list(linha) for linha in planilha.iter_rows(values_only=True)]
    else:
        with open(caminho, encoding="utf-8-sig", newline="") as arquivo:
            amostra = arquivo.read(4096)
            arquivo.seek(0)
            separador = ";" if amostra.count(";") > amostra.count(",") else ","
            linhas = list(csv.reader(arquivo, delimiter=separador))

    cabecalhos, dados = linhas[0], linhas[1:]
    posicoes, usadas = {}, set()
    for campo in ("subsetor", "setor", "cnpj", "nome", "telefone", "site", "faturamento"):
        indice = _achar_coluna(cabecalhos, MAPA_COLUNAS[campo], usadas)
        posicoes[campo] = indice
        if indice is not None:
            usadas.add(indice)

    empresas = []
    for linha in dados:
        if not any(linha):
            continue
        empresa = {
            campo: (linha[i] if i is not None and i < len(linha) else "")
            for campo, i in posicoes.items()
        }
        empresas.append(empresa)
    return empresas


def _texto(valor):
    return "" if valor is None else str(valor).strip()


def carregar_lista_driva(caminho):
    """Lê o export do Driva (várias abas) e junta tudo por CNPJ:
    aba Empresa + telefones fixos (aba Telefones) + sites (aba Sites) + redes sociais."""
    pasta = load_workbook(caminho, read_only=True)

    def ler_aba(nome):
        if nome not in pasta.sheetnames:
            return []
        linhas = list(pasta[nome].iter_rows(values_only=True))
        if not linhas:
            return []
        cabecalho = [_texto(c) for c in linhas[0]]
        return [dict(zip(cabecalho, linha)) for linha in linhas[1:] if any(linha)]

    telefones, sites = {}, {}
    for t in ler_aba("Telefones"):
        if _texto(t.get("Fixo Ou Móvel")).lower() == "fixo":
            telefones.setdefault(limpar_cnpj(t.get("CNPJ")), []).append(_texto(t.get("Telefone Completo")))
    for s in ler_aba("Sites"):
        if _texto(s.get("Pertence ao Contador")).lower() != "sim" and _texto(s.get("URL")):
            sites.setdefault(limpar_cnpj(s.get("CNPJ")), []).append(_texto(s.get("URL")))

    empresas = []
    for e in ler_aba("Empresa"):
        cnpj = limpar_cnpj(e.get("CNPJ"))
        lista_sites = sites.get(cnpj, [])
        empresas.append({
            "origem": "Driva",
            "nome": _texto(e.get("Razão Social")),
            "nome_fantasia": _texto(e.get("Nome Fantasia")),
            "cnpj": cnpj,
            "telefone": "; ".join(extrair_telefones_fixos(" ; ".join(telefones.get(cnpj, [])))),
            "site": lista_sites[0] if len(lista_sites) == 1 else "",
            "sites_candidatos": lista_sites,
            "setor": _texto(e.get("Setor")),
            "subsetor": _texto(e.get("Subsetor")),
            "faturamento": _texto(e.get("Faixa de Faturamento Grupo")),
            "atividade": _texto(e.get("CNAE Principal Desc. Subclasse")),
            "cidade": f"{_texto(e.get('Município'))}/{_texto(e.get('UF'))}",
            "linkedin": _texto(e.get("Linkedin Melhor Resultado")),
            "facebook": _texto(e.get("Facebook Melhor Resultado")),
            "instagram": _texto(e.get("Instagram Melhor Resultado")),
        })
    return empresas


def carregar_lista(caminho):
    """Detecta a origem do arquivo (Driva ou Speedio/outro) e lê do jeito certo."""
    if caminho.lower().endswith(".xlsx") and "Empresa" in load_workbook(caminho, read_only=True).sheetnames:
        return carregar_lista_driva(caminho)
    empresas = carregar_lista_speedio(caminho)
    for e in empresas:
        e["origem"] = "Speedio"
    return empresas


# ----------------------------------------------------------------------
# 6. Gerar a planilha final
# ----------------------------------------------------------------------
COLUNAS_PLANILHA = [
    "Nome", "CNPJ", "Telefone fixo (lista de origem)", "Site", "Setor", "Subsetor",
    "Faturamento anual", "Telefones encontrados", "Fontes dos telefones",
    "Dentro do perfil?", "Observação",
]


def salvar_planilha(empresas, caminho):
    """Salva a lista de empresas no .xlsx. Os telefones ficam separados por ';'."""
    pasta = Workbook()
    aba = pasta.active
    aba.title = "Leads"
    aba.append(COLUNAS_PLANILHA)

    for celula in aba[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F4E78")

    for e in empresas:
        aba.append([
            e.get("nome", ""),
            e.get("cnpj", ""),
            e.get("telefone", ""),
            e.get("site", ""),
            e.get("setor", ""),
            e.get("subsetor", ""),
            e.get("faturamento", ""),
            "; ".join(e.get("telefones_encontrados", [])),
            "; ".join(e.get("fontes", [])),
            "Sim" if e.get("dentro_do_perfil", True) else "CONFERIR",
            e.get("observacao", ""),
        ])

    larguras = [40, 20, 22, 30, 25, 30, 20, 50, 60, 16, 50]
    for coluna, largura in zip("ABCDEFGHIJK", larguras):
        aba.column_dimensions[coluna].width = largura
    aba.freeze_panes = "A2"

    pasta.save(caminho)
    return caminho


# ----------------------------------------------------------------------
# Teste rápido de todas as ferramentas
# ----------------------------------------------------------------------
if __name__ == "__main__":
    print("1) Teste de telefones:")
    print("  ", extrair_telefones_fixos("Ligue (19) 3577-9705, cel (19) 99999-1234 ou 0800 123 4567"))

    print("\n2) Receita Federal:")
    receita = consultar_cnpj("00.000.000/0001-91")
    print("  ", receita)

    print("\n3) Busca no Google:")
    busca = buscar_google(f"{receita.get('razao_social', 'Banco do Brasil')} telefone")
    if "erro" in busca:
        print("  ", busca["erro"])
    else:
        for r in busca["resultados"][:3]:
            print("   -", r["titulo"], "|", r["link"], "|", r["telefones_no_resumo"])

    print("\n4) Leitura de site:")
    site = ler_site("https://www.bb.com.br")
    print("   Telefones:", site.get("telefones_fixos"), "| Erro:", site.get("erro"))
    print("   Links de contato:", site.get("links_de_contato", [])[:3])

    print("\n5) Planilha de teste:")
    exemplo = [{
        "nome": receita.get("razao_social"), "cnpj": receita.get("cnpj"),
        "telefones_encontrados": receita.get("telefones_fixos", []),
        "fontes": ["Receita Federal"],
    }]
    print("   Salva em:", salvar_planilha(exemplo, "teste_planilha.xlsx"))
